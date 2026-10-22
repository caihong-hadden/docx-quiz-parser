import os
import re
import json
import docx

# Directories
converted_dir = r"C:\Users\SunnyRainbow\.gemini\antigravity\scratch\maogai-quiz\converted"
output_json_path = r"C:\Users\SunnyRainbow\.gemini\antigravity\scratch\maogai-quiz\questions.json"

def clean_text(text):
    # Remove leading/trailing whitespaces and special spaces
    text = text.strip()
    text = text.replace('\xa0', ' ') # non-breaking space
    text = text.replace('\u3000', ' ') # ideographic space
    return text

def split_options(line):
    # Splitting A. B. C. D.
    # Uses a lookahead to split right before an option prefix
    pattern = r'(?=[A-D][\.．\s])'
    parts = re.split(pattern, line)
    options = [clean_text(p) for p in parts if p.strip()]
    return options

def parse_embedded_answer(stem):
    # Matches patterns like （ B ）, (ABCD), （  C  ）, etc.
    # Inside parentheses, we look for letters A, B, C, D
    pattern = r'[（(]\s*([A-D\s]+)\s*[）)]'
    match = re.search(pattern, stem)
    if match:
        answer = match.group(1).replace(" ", "").strip()
        # Replace the answer with blanks
        cleaned_stem = re.sub(pattern, '（   ）', stem)
        return cleaned_stem, answer
    return stem, None

def extract_table_answers(table):
    """
    Extracts mappings of {question_num: answer} from a Word table.
    We look for pairs of rows where row K has numbers and row K+1 has answers.
    """
    answers = {}
    for k in range(0, len(table.rows) - 1, 2):
        row_q = table.rows[k]
        row_a = table.rows[k+1]
        
        q_cells = [clean_text(cell.text) for cell in row_q.cells]
        a_cells = [clean_text(cell.text) for cell in row_a.cells]
        
        # Check if the row contains headers/labels
        for q_val, a_val in zip(q_cells, a_cells):
            # Check if q_val is a number
            if q_val.isdigit():
                num = int(q_val)
                # Clean answer text (e.g. remove spaces, dots)
                ans = a_val.replace(" ", "").replace(".", "").strip().upper()
                if ans:
                    answers[num] = ans
    return answers

def extract_text_answers(text):
    """
    Extracts answers from plain text block like "1.B  2.A  3.C" or "1.ABCD  2.ABCD".
    """
    answers = {}
    # Match patterns like: 1.B  or 1.ABCD or 1. A or 1．B
    pattern = r'(\d+)\s*[\.．]\s*([A-D\s]+)'
    matches = re.findall(pattern, text)
    for q_num, ans in matches:
        num = int(q_num)
        ans_clean = ans.replace(" ", "").strip().upper()
        if ans_clean:
            answers[num] = ans_clean
    return answers

def parse_docx_file(filepath, file_id):
    doc = docx.Document(filepath)
    paragraphs = doc.paragraphs
    tables = doc.tables
    
    # Check if this is the Chapter 1 file which needs to be split into "Intro" and "Chapter 1"
    is_ch1_file = "导论和第1章" in filepath or "chapter1.docx" in filepath or "导论" in os.path.basename(filepath)
    
    if is_ch1_file:
        return parse_ch1_combined_file(paragraphs, tables)
    else:
        return parse_standard_file(paragraphs, tables, file_id)

def parse_ch1_combined_file(paragraphs, tables):
    # Introduction + Chapter 1
    # We will create two separate lists of questions
    intro_questions = []
    ch1_questions = []
    
    current_section = "intro" # "intro" or "ch1"
    current_q_type = "single" # "single", "multiple", "short", "essay"
    
    # Store questions for processing
    intro_raw_questions = []
    ch1_raw_questions = []
    
    # Parse paragraphs
    for p_idx, p in enumerate(paragraphs):
        text = clean_text(p.text)
        if not text:
            continue
            
        # Detect section switches
        if "第一章" in text:
            current_section = "ch1"
            current_q_type = None
            continue
        elif "导论" in text:
            current_section = "intro"
            current_q_type = "single"
            continue
            
        # Detect question types
        if "单项选择题" in text or "单选题" in text:
            current_q_type = "single"
            continue
        elif "多项选择题" in text or "多选题" in text:
            current_q_type = "multiple"
            continue
        elif "简答题" in text:
            current_q_type = "short"
            continue
        elif "论述题" in text:
            current_q_type = "essay"
            continue
            
        # Skip reference answer labels in text stream
        if "参考答案" in text or "答案要点" in text:
            continue
            
        # Parse question items
        if current_q_type in ["single", "multiple"]:
            # Check if this paragraph starts a new question (e.g. "1.", "10.")
            match_q = re.match(r'^(\d+)[\.．\s]', text)
            if match_q:
                q_num = int(match_q.group(1))
                stem = text
                options = []
                
                # Split options if option A is embedded in the stem paragraph
                if "A." in stem or "A．" in stem:
                    idx_a = stem.find("A.") if "A." in stem else stem.find("A．")
                    stem_part = stem[:idx_a].strip()
                    opt_part = stem[idx_a:].strip()
                    
                    stem_clean, embedded_ans = parse_embedded_answer(stem_part)
                    split_opts = split_options(opt_part)
                    options.extend(split_opts)
                else:
                    stem_clean, embedded_ans = parse_embedded_answer(stem)
                
                # Gather options
                lookahead_idx = p_idx + 1
                while lookahead_idx < len(paragraphs):
                    next_text = clean_text(paragraphs[lookahead_idx].text)
                    if not next_text:
                        lookahead_idx += 1
                        continue
                    # If it starts a new question, or is a section heading, or is answers block, stop
                    if re.match(r'^(\d+)[\.．\s]', next_text) or any(h in next_text for h in ["单选", "多选", "简答", "论述", "参考答案", "第一章", "答案"]):
                        break
                    # Check if it starts with an option prefix (e.g. A., B.)
                    has_prefix = any(next_text.startswith(x) for x in ["A.", "B.", "C.", "D.", "A．", "B．", "C．", "D．"])
                    if has_prefix:
                        split_opts = split_options(next_text)
                        options.extend(split_opts)
                    else:
                        # Auto prepend letter prefix!
                        next_letter = chr(ord('A') + len(options))
                        options.append(f"{next_letter}. {next_text}")
                    lookahead_idx += 1
                    
                raw_item = {
                    "num": q_num,
                    "stem": stem_clean,
                    "options": options,
                    "answer": embedded_ans,
                    "type": current_q_type
                }
                
                if current_section == "intro":
                    intro_raw_questions.append(raw_item)
                else:
                    ch1_raw_questions.append(raw_item)
                    
        elif current_q_type in ["short", "essay"]:
            # Parse short/essay question stem
            match_q = re.match(r'^(\d+)[\.．\s]', text)
            if match_q:
                q_num = int(match_q.group(1))
                stem = text
                
                # Check for inline answers immediately following the subjective question
                inline_ans_parts = []
                lookahead_idx = p_idx + 1
                while lookahead_idx < len(paragraphs):
                    next_text = clean_text(paragraphs[lookahead_idx].text)
                    if not next_text:
                        lookahead_idx += 1
                        continue
                    # Stop if we hit a new question or section heading
                    if re.match(r'^(\d+)[\.．\s]', next_text) or any(h in next_text for h in ["单选题", "多选题", "简答题", "论述题", "一、", "二、", "三、", "四、"]):
                        if "参考答案" in next_text or "答案要点" in next_text:
                            inline_ans_parts.append(next_text)
                            lookahead_idx += 1
                            continue
                        break
                    inline_ans_parts.append(next_text)
                    lookahead_idx += 1
                
                ans_str = ""
                if inline_ans_parts:
                    ans_str = "\n".join(inline_ans_parts).strip()
                    ans_str = re.sub(r'^(答：|答案要点：|参考答案：|答\s*:\s*)', '', ans_str).strip()
                
                raw_item = {
                    "num": q_num,
                    "stem": stem,
                    "type": current_q_type,
                    "answer": ans_str
                }
                if current_section == "intro":
                    intro_raw_questions.append(raw_item)
                else:
                    ch1_raw_questions.append(raw_item)

    # Now let's extract answers for Chapter 1 and Introduction
    # Introduction has answers embedded. If any intro question has no answer, let's verify if there is an answer block
    # Chapter 1 answers are in tables.
    ch1_single_answers = {}
    ch1_multiple_answers = {}
    
    if len(tables) >= 2:
        ch1_single_answers = extract_table_answers(tables[0])
        ch1_multiple_answers = extract_table_answers(tables[1])
        print(f"Ch1: Extracted {len(ch1_single_answers)} single choice and {len(ch1_multiple_answers)} multiple choice answers from tables.")

    # Let's extract short answers/essays answers from paragraphs
    # We will look for "参考答案" sections in paragraphs
    intro_short_answers = {}
    ch1_short_answers = {}
    ch1_essay_answers = {}
    
    curr_ans_section = None # "intro_short", "ch1_short", "ch1_essay"
    curr_ans_text = []
    curr_ans_num = None
    
    for idx, p in enumerate(paragraphs):
        text = clean_text(p.text)
        if not text:
            continue
            
        # Stop answer parsing if we exit the answer block and hit a new section heading
        if curr_ans_section and any(h in text for h in ["一、单项选择题", "二、多项选择题", "第一章", "第一章 （参考答案）"]):
            if curr_ans_num is not None:
                ans_str = "\n".join(curr_ans_text).strip()
                ans_str = re.sub(r'^(答：|答案要点：|参考答案：|答\s*:\s*)', '', ans_str).strip()
                if curr_ans_section == "intro_short":
                    intro_short_answers[curr_ans_num] = ans_str
                elif curr_ans_section == "ch1_short":
                    ch1_short_answers[curr_ans_num] = ans_str
                elif curr_ans_section == "ch1_essay":
                    ch1_essay_answers[curr_ans_num] = ans_str
            curr_ans_section = None
            curr_ans_num = None
            curr_ans_text = []
            continue
            
        # Detect answer headings
        if "三、简答题" in text:
            # Save previous answer first
            if curr_ans_num is not None and curr_ans_section:
                ans_str = "\n".join(curr_ans_text).strip()
                ans_str = re.sub(r'^(答：|答案要点：|参考答案：|答\s*:\s*)', '', ans_str).strip()
                if curr_ans_section == "intro_short":
                    intro_short_answers[curr_ans_num] = ans_str
                elif curr_ans_section == "ch1_short":
                    ch1_short_answers[curr_ans_num] = ans_str
                elif curr_ans_section == "ch1_essay":
                    ch1_essay_answers[curr_ans_num] = ans_str
            
            if "第一章" in text or idx > 150: # Ch1 answers are later in the doc
                curr_ans_section = "ch1_short"
            else:
                curr_ans_section = "intro_short"
            curr_ans_num = None
            curr_ans_text = []
            continue
        elif "四、论述题" in text:
            # Save previous answer first
            if curr_ans_num is not None and curr_ans_section:
                ans_str = "\n".join(curr_ans_text).strip()
                ans_str = re.sub(r'^(答：|答案要点：|参考答案：|答\s*:\s*)', '', ans_str).strip()
                if curr_ans_section == "intro_short":
                    intro_short_answers[curr_ans_num] = ans_str
                elif curr_ans_section == "ch1_short":
                    ch1_short_answers[curr_ans_num] = ans_str
                elif curr_ans_section == "ch1_essay":
                    ch1_essay_answers[curr_ans_num] = ans_str
            
            curr_ans_section = "ch1_essay"
            curr_ans_num = None
            curr_ans_text = []
            continue
        elif "第一章 （参考答案）" in text or "第三章" in text or "参考答案" in text:
            # Skip overall headers
            continue
            
        # If we are in an answer section
        if curr_ans_section in ["intro_short", "ch1_short", "ch1_essay"]:
            # Check if this paragraph starts a new answer (e.g. "1.", "2.")
            match_num = re.match(r'^(\d+)[\.．\s]', text)
            if match_num:
                # Save previous answer
                if curr_ans_num is not None:
                    ans_str = "\n".join(curr_ans_text).strip()
                    # Clean answer prefix like "答：" or "答案要点："
                    ans_str = re.sub(r'^(答：|答案要点：|参考答案：|答\s*:\s*)', '', ans_str).strip()
                    if curr_ans_section == "intro_short":
                        intro_short_answers[curr_ans_num] = ans_str
                    elif curr_ans_section == "ch1_short":
                        ch1_short_answers[curr_ans_num] = ans_str
                    elif curr_ans_section == "ch1_essay":
                        ch1_essay_answers[curr_ans_num] = ans_str
                
                curr_ans_num = int(match_num.group(1))
                curr_ans_text = [text]
            else:
                if curr_ans_num is not None:
                    curr_ans_text.append(text)
                    
    # Save the last active short/essay answer
    if curr_ans_num is not None:
        ans_str = "\n".join(curr_ans_text).strip()
        ans_str = re.sub(r'^(答：|答案要点：|参考答案：|答\s*:\s*)', '', ans_str).strip()
        if curr_ans_section == "intro_short":
            intro_short_answers[curr_ans_num] = ans_str
        elif curr_ans_section == "ch1_short":
            ch1_short_answers[curr_ans_num] = ans_str
        elif curr_ans_section == "ch1_essay":
            ch1_essay_answers[curr_ans_num] = ans_str

    # Process and compile intro questions
    final_intro = []
    for item in intro_raw_questions:
        q_id = f"intro_{item['type']}_{item['num']}"
        ans = item['answer']
        if item['type'] == 'short' and item['num'] in intro_short_answers:
            ans = intro_short_answers[item['num']]
            
        final_intro.append({
            "id": q_id,
            "type": item['type'],
            "stem": item['stem'],
            "options": item.get('options', []),
            "answer": ans,
            "explanation": ""
        })
        
    # Process and compile ch1 questions
    final_ch1 = []
    for item in ch1_raw_questions:
        q_id = f"ch1_{item['type']}_{item['num']}"
        ans = item['answer']
        if item['type'] == 'single':
            ans = ch1_single_answers.get(item['num'], ans)
        elif item['type'] == 'multiple':
            ans = ch1_multiple_answers.get(item['num'], ans)
        elif item['type'] == 'short':
            ans = ch1_short_answers.get(item['num'], ans)
        elif item['type'] == 'essay':
            ans = ch1_essay_answers.get(item['num'], ans)
            
        final_ch1.append({
            "id": q_id,
            "type": item['type'],
            "stem": item['stem'],
            "options": item.get('options', []),
            "answer": ans,
            "explanation": ""
        })
        
    return {
        "intro": {"name": "导论 马克思主义中国化时代化的历史进程与理论成果", "questions": final_intro},
        "ch1": {"name": "第一章 毛泽东思想及其历史地位", "questions": final_ch1}
    }

def parse_standard_file(paragraphs, tables, chapter_id):
    current_q_type = None
    raw_questions = []
    
    # Store plain text answers block
    answers_text_block = []
    in_answers_block = False
    
    # Main paragraph pass
    for p_idx, p in enumerate(paragraphs):
        text = clean_text(p.text)
        if not text:
            continue
            
        if in_answers_block:
            answers_text_block.append(text)
            continue
            
        # Detect question types
        if "单项选择题" in text or "单选题" in text:
            current_q_type = "single"
            continue
        elif "多项选择题" in text or "多选题" in text:
            current_q_type = "multiple"
            continue
        elif "简答题" in text:
            current_q_type = "short"
            continue
        elif "论述题" in text:
            current_q_type = "essay"
            continue
        elif "参考答案" in text or "答案要点" in text or "答案（" in text:
            # Switch to answers parsing block
            current_q_type = "answers_block"
            in_answers_block = True
            continue
            
        # Parse question items
        if current_q_type in ["single", "multiple"]:
            match_q = re.match(r'^(\d+)[\.．\s]', text)
            if match_q:
                q_num = int(match_q.group(1))
                stem = text
                options = []
                
                # Split options if option A is embedded in the stem paragraph
                if "A." in stem or "A．" in stem:
                    idx_a = stem.find("A.") if "A." in stem else stem.find("A．")
                    stem_part = stem[:idx_a].strip()
                    opt_part = stem[idx_a:].strip()
                    
                    stem_clean, embedded_ans = parse_embedded_answer(stem_part)
                    split_opts = split_options(opt_part)
                    options.extend(split_opts)
                else:
                    stem_clean, embedded_ans = parse_embedded_answer(stem)
                
                # Gather options
                lookahead_idx = p_idx + 1
                inline_ans = None
                while lookahead_idx < len(paragraphs):
                    next_text = clean_text(paragraphs[lookahead_idx].text)
                    if not next_text:
                        lookahead_idx += 1
                        continue
                    if re.match(r'^(\d+)[\.．\s]', next_text) or any(h in next_text for h in ["单选", "多选", "简答", "论述", "参考答案", "答案"]):
                        # Check if next line is the answer (Chapter 5 format: e.g. "答案B")
                        if next_text.startswith("答案"):
                            ans_match = re.match(r'^答案[\.．\s]?\s*([A-D]+)', next_text)
                            if ans_match:
                                inline_ans = ans_match.group(1).strip().upper()
                        break
                    # Check if it starts with an option prefix (e.g. A., B.)
                    has_prefix = any(next_text.startswith(x) for x in ["A.", "B.", "C.", "D.", "A．", "B．", "C．", "D．"])
                    if has_prefix:
                        split_opts = split_options(next_text)
                        options.extend(split_opts)
                    else:
                        # Auto prepend letter prefix!
                        next_letter = chr(ord('A') + len(options))
                        options.append(f"{next_letter}. {next_text}")
                    lookahead_idx += 1
                    
                raw_item = {
                    "num": q_num,
                    "stem": stem_clean,
                    "options": options,
                    "answer": embedded_ans or inline_ans,
                    "type": current_q_type
                }
                raw_questions.append(raw_item)
                
        elif current_q_type in ["short", "essay"]:
            match_q = re.match(r'^(\d+)[\.．\s]', text)
            if match_q:
                q_num = int(match_q.group(1))
                stem = text
                
                # Check for inline answers immediately following the subjective question
                inline_ans_parts = []
                lookahead_idx = p_idx + 1
                while lookahead_idx < len(paragraphs):
                    next_text = clean_text(paragraphs[lookahead_idx].text)
                    if not next_text:
                        lookahead_idx += 1
                        continue
                    # Stop if we hit a new question or section heading
                    if re.match(r'^(\d+)[\.．\s]', next_text) or any(h in next_text for h in ["单选题", "多选题", "简答题", "论述题", "一、", "二、", "三、", "四、"]):
                        if "参考答案" in next_text or "答案要点" in next_text:
                            inline_ans_parts.append(next_text)
                            lookahead_idx += 1
                            continue
                        break
                    inline_ans_parts.append(next_text)
                    lookahead_idx += 1
                
                ans_str = ""
                if inline_ans_parts:
                    ans_str = "\n".join(inline_ans_parts).strip()
                    ans_str = re.sub(r'^(答：|答案要点：|参考答案：|答\s*:\s*)', '', ans_str).strip()
                
                raw_item = {
                    "num": q_num,
                    "stem": stem,
                    "type": current_q_type,
                    "answer": ans_str
                }
                raw_questions.append(raw_item)

    # Resolve Answers
    single_answers = {}
    multiple_answers = {}
    short_essay_answers = {}
    
    # 1. Check Tables (Chapters 2, 6, 7, 8)
    if len(tables) >= 2:
        single_answers = extract_table_answers(tables[0])
        multiple_answers = extract_table_answers(tables[1])
        print(f"{chapter_id}: Extracted {len(single_answers)} single and {len(multiple_answers)} multiple choice answers from tables.")
    
    # 2. Check Answers Text Block (Chapters 3, 4, 5)
    # Parse plain text single/multiple choice answers
    text_answers_str = " ".join(answers_text_block)
    
    # Split text answers block by headings to distinguish single and multiple choice
    single_text_block = ""
    multiple_text_block = ""
    short_essay_text_block = []
    
    curr_block_type = None
    for line in answers_text_block:
        if "单选题" in line or "单项选择题" in line:
            curr_block_type = "single"
            continue
        elif "多选题" in line or "多项选择题" in line:
            curr_block_type = "multiple"
            continue
        elif "简答题" in line or "论述题" in line:
            curr_block_type = "short_essay"
            continue
            
        if curr_block_type == "single":
            single_text_block += " " + line
        elif curr_block_type == "multiple":
            multiple_text_block += " " + line
        elif curr_block_type == "short_essay":
            short_essay_text_block.append(line)
            
    # Extract
    if single_text_block:
        single_answers.update(extract_text_answers(single_text_block))
    else:
        # If no explicit headings, try extract all from the text answers block
        single_answers.update(extract_text_answers(text_answers_str))
        
    if multiple_text_block:
        multiple_answers.update(extract_text_answers(multiple_text_block))
        
    # Extract short answers/essays
    # Parse list of Q & A from short_essay_text_block or overall answers_text_block
    target_block = short_essay_text_block if short_essay_text_block else answers_text_block
    curr_q_num = None
    curr_q_text = []
    
    for line in target_block:
        match_num = re.match(r'^(\d+)[\.．\s]', line)
        if match_num:
            if curr_q_num is not None:
                ans_str = "\n".join(curr_q_text).strip()
                ans_str = re.sub(r'^(答：|答案要点：|参考答案：|答\s*:\s*)', '', ans_str).strip()
                short_essay_answers[curr_q_num] = ans_str
            curr_q_num = int(match_num.group(1))
            curr_q_text = [line]
        else:
            if curr_q_num is not None:
                curr_q_text.append(line)
                
    if curr_q_num is not None:
        ans_str = "\n".join(curr_q_text).strip()
        ans_str = re.sub(r'^(答：|答案要点：|参考答案：|答\s*:\s*)', '', ans_str).strip()
        short_essay_answers[curr_q_num] = ans_str

    # Process and compile final questions list
    final_questions = []
    for item in raw_questions:
        q_id = f"{chapter_id}_{item['type']}_{item['num']}"
        ans = item['answer']
        if item['type'] == 'single':
            ans = single_answers.get(item['num'], ans)
        elif item['type'] == 'multiple':
            ans = multiple_answers.get(item['num'], ans)
        elif item['type'] in ['short', 'essay']:
            ans = short_essay_answers.get(item['num'], ans)
            
        # Fallback if ans is still None or empty (for single/multiple choice)
        if not ans and item['type'] in ['single', 'multiple']:
            ans = ""
            
        # Fallback for Chapter 8 Q16
        if chapter_id == "ch8" and item['type'] == "single" and item['num'] == 16 and not ans:
            ans = "B"
            
        final_questions.append({
            "id": q_id,
            "type": item['type'],
            "stem": item['stem'],
            "options": item.get('options', []),
            "answer": ans,
            "explanation": ""
        })
        
    return final_questions

def main():
    files_map = [
        {"id": "ch1", "file": "导论和第1章毛泽东思想及其历史地位习题及答案.docx"},
        {"id": "ch2", "file": "第2章新民主主义革命理论习题及答案.docx"},
        {"id": "ch3", "file": "第3章社会主义改造理论习题及答案.docx"},
        {"id": "ch4", "file": "第4章社会主义建设道路初步探索的理论成果习题及答案.docx"},
        {"id": "ch5", "file": "第5章中国特色社会主义理论体系的形成发展.docx"},
        {"id": "ch6", "file": "第6章邓小平理论.docx"},
        {"id": "ch7", "file": "第7章 “三个代表”重要思想 习题和答案.docx"},
        {"id": "ch8", "file": "第8章科学发展观习题和答案.docx"},
    ]
    
    database = {}
    
    for item in files_map:
        filepath = os.path.join(converted_dir, item["file"])
        if not os.path.exists(filepath):
            print(f"File not found: {filepath}")
            continue
            
        print(f"Parsing {item['file']}...")
        if item["id"] == "ch1":
            # This file returns two chapters: intro and ch1
            results = parse_docx_file(filepath, "ch1")
            database["intro"] = results["intro"]
            database["ch1"] = results["ch1"]
            print(f"  Intro: {len(results['intro']['questions'])} questions")
            print(f"  Ch1: {len(results['ch1']['questions'])} questions")
        else:
            questions = parse_docx_file(filepath, item["id"])
            
            # Map chapter IDs to textbook names
            chapter_names = {
                "ch2": "第二章 新民主主义革命理论",
                "ch3": "第三章 社会主义改造理论",
                "ch4": "第四章 社会主义建设道路初步探索的理论成果",
                "ch5": "第五章 中国特色社会主义理论体系的形成发展",
                "ch6": "第六章 邓小平理论",
                "ch7": "第七章 “三个代表”重要思想",
                "ch8": "第八章 科学发展观",
            }
            
            database[item["id"]] = {
                "name": chapter_names.get(item["id"], item["id"]),
                "questions": questions
            }
            print(f"  {item['id']}: {len(questions)} questions")
            
    # Save database to JSON
    with open(output_json_path, "w", encoding="utf-8") as f:
        json.dump(database, f, ensure_ascii=False, indent=2)
    print(f"\nSuccessfully compiled questions database! Saved to {output_json_path}")
    
    # Print statistics
    total_q = 0
    for ch_id, ch_data in database.items():
        types_count = {}
        for q in ch_data["questions"]:
            types_count[q["type"]] = types_count.get(q["type"], 0) + 1
            total_q += 1
        print(f"{ch_id} ({ch_data['name']}): {len(ch_data['questions'])} questions | {types_count}")
    print(f"Total questions compiled: {total_q}")

if __name__ == "__main__":
    main()
