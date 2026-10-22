import docx

docx_path = r"C:\Users\SunnyRainbow\.gemini\antigravity\scratch\maogai-quiz\converted\chapter1.docx"
doc = docx.Document(docx_path)

print("Table count:", len(doc.tables))
for index, table in enumerate(doc.tables):
    print(f"\n--- Table {index} ---")
    for r_idx, row in enumerate(table.rows):
        row_text = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        print(f"Row {r_idx}: {row_text}")
