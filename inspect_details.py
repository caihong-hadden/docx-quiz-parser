import docx
import os

converted_dir = r"C:\Users\SunnyRainbow\.gemini\antigravity\scratch\maogai-quiz\converted"

# Inspect Chapter 4 answers
doc4 = docx.Document(os.path.join(converted_dir, "第4章社会主义建设道路初步探索的理论成果习题及答案.docx"))
print("=== Chapter 4 Answers ===")
ans_start = False
for idx, p in enumerate(doc4.paragraphs):
    text = p.text.strip()
    if "参考答案" in text or ans_start:
        ans_start = True
        print(f"{idx}: {text}")

# Inspect Chapter 8 answers
doc8 = docx.Document(os.path.join(converted_dir, "第8章科学发展观习题和答案.docx"))
print("\n=== Chapter 8 Answers ===")
ans_start = False
for idx, p in enumerate(doc8.paragraphs):
    text = p.text.strip()
    if "参考答案" in text or ans_start:
        ans_start = True
        print(f"{idx}: {text}")
