import docx
import os

converted_dir = r"C:\Users\SunnyRainbow\.gemini\antigravity\scratch\maogai-quiz\converted"
doc = docx.Document(os.path.join(converted_dir, "第5章中国特色社会主义理论体系的形成发展.docx"))

for idx in range(min(len(doc.paragraphs), 60), min(len(doc.paragraphs), 98)):
    print(f"{idx}: {doc.paragraphs[idx].text}")
