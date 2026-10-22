import docx
import os

converted_dir = r"C:\Users\SunnyRainbow\.gemini\antigravity\scratch\maogai-quiz\converted"
doc = docx.Document(os.path.join(converted_dir, "第3章社会主义改造理论习题及答案.docx"))

print("Total paragraphs:", len(doc.paragraphs))
# Print the last 40 paragraphs
for i in range(len(doc.paragraphs) - 40, len(doc.paragraphs)):
    print(f"{i}: {doc.paragraphs[i].text}")
