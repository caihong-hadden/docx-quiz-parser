import docx
import os

converted_dir = r"C:\Users\SunnyRainbow\.gemini\antigravity\scratch\maogai-quiz\converted"
doc = docx.Document(os.path.join(converted_dir, "第5章中国特色社会主义理论体系的形成发展.docx"))

with open("ch5_end_paragraphs.txt", "w", encoding="utf-8") as f:
    for idx in range(92, len(doc.paragraphs)):
        f.write(f"Index {idx}: {repr(doc.paragraphs[idx].text)}\n")
print("Written successfully!")
