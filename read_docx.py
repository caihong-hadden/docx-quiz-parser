import docx
import os

docx_path = r"E:\Desktop\2026毛概论题库\2026毛概论题库\第3章社会主义改造理论习题及答案.docx"
doc = docx.Document(docx_path)
with open("docx_content.txt", "w", encoding="utf-8") as f:
    f.write(f"Paragraph count: {len(doc.paragraphs)}\n")
    for i, p in enumerate(doc.paragraphs):
        f.write(f"{i}: {p.text}\n")
print("Written successfully to docx_content.txt")
